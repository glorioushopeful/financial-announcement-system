# -*- coding: utf-8 -*-
"""导出模块：把清洗后的 DataFrame 落盘为 Excel / CSV。

Excel 带表头样式、冻结首行、自动列宽与筛选器，可直接交付业务方查看。
"""

import os
from typing import Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

# 中文表头映射，让 Excel 可读性更好
HEADER_CN = {
    "stock_code": "股票代码",
    "stock_name": "股票名称",
    "title": "公告标题",
    "category": "公告分类",
    "publish_date": "发布日期",
    "year": "年份",
    "month": "月份",
    "quarter": "季度",
    "weekday": "星期(0=周一)",
    "doc_type": "文档类型",
    "title_length": "标题长度",
    "url": "原文链接",
    "source": "数据来源",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def to_csv(df: pd.DataFrame, path: str = None) -> str:
    """导出为 UTF-8-BOM CSV，保证 Excel 直接打开不乱码。"""
    path = path or config.CSV_PATH
    os.makedirs(os.path.dirname(path), exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def to_excel(df: pd.DataFrame, path: str = None) -> str:
    """导出为带格式的 Excel。"""
    path = path or config.EXCEL_PATH
    os.makedirs(os.path.dirname(path), exist_ok=True)

    out = df.rename(columns=HEADER_CN)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        out.to_excel(writer, index=False, sheet_name="公告数据")
        _beautify(writer.sheets["公告数据"], out)

    return path


def _beautify(ws, df: pd.DataFrame) -> None:
    """设置表头样式、自动列宽、冻结首行与筛选。"""
    # 表头
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 24

    # 自动列宽（按内容长度估算，中文按 2 个字符宽计算）
    for i, col in enumerate(df.columns, start=1):
        sample = df[col].astype(str).head(200)
        max_len = max([len(str(col))] + [len(s) * 2 if any("\u4e00" <= c <= "\u9fff" for c in s)
                                         else len(s) for s in sample]) if len(sample) else len(str(col))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 4, 10), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export(df: pd.DataFrame, logger=print) -> Tuple[str, str]:
    """同时导出 CSV 与 Excel。"""
    csv_path = to_csv(df)
    xlsx_path = to_excel(df)
    logger(f"  已导出 CSV : {csv_path}")
    logger(f"  已导出 Excel: {xlsx_path}")
    return csv_path, xlsx_path
